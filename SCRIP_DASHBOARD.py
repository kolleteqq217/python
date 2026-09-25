import io
import streamlit as st
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from google.oauth2 import service_account
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload, MediaIoBaseUpload

# --- Configuración visual ---
st.set_page_config(
    page_title="Dashboard Comercial | GlobalTech",
    page_icon="📊",
    layout="wide"
)

# --- Función de limpieza numérica ---
def clean_val(v):
    if pd.isna(v) or v is None:
        return 0.0
    s = str(v).strip().replace('$', '').replace(' ', '').replace('\xa0', '')
    if not s or s.lower() in ('nan', 'none', 'null'):
        return 0.0
    if '.' in s and ',' in s:
        if s.rfind('.') > s.rfind(','):
            s = s.replace(',', '')
        else:
            s = s.replace('.', '').replace(',', '.')
    elif ',' in s:
        parts = s.split(',')
        if len(parts) == 2 and len(parts[1]) in (1, 2):
            s = s.replace(',', '.')
        else:
            s = s.replace(',', '')
    try:
        return float(s)
    except Exception:
        return 0.0

# --- 1. Autenticación con Google Drive ---
SCOPES = ['https://www.googleapis.com/auth/drive']

@st.cache_resource
def get_drive_service():
    """Inicializa el cliente de Google Drive."""
    creds_dict = dict(st.secrets["gcp_service_account"])
    creds = service_account.Credentials.from_service_account_info(creds_dict, scopes=SCOPES)
    return build('drive', 'v3', credentials=creds)

# --- 2. Descargar datos desde Google Drive ---
@st.cache_data(ttl=60)
def load_data_from_drive(file_id):
    """Descarga el Excel en memoria."""
    service = get_drive_service()
    request = service.files().get_media(fileId=file_id, supportsAllDrives=True)
    fh = io.BytesIO()
    downloader = MediaIoBaseDownload(fh, request)
    done = False
    while not done:
        _, done = downloader.next_chunk()
    fh.seek(0)
    
    # Cargar pestaña 'Ventas'
    try:
        df = pd.read_excel(fh, sheet_name="Ventas", engine='openpyxl')
    except Exception:
        fh.seek(0)
        df = pd.read_excel(fh, sheet_name=0, engine='openpyxl')
        
    df.columns = [str(c).strip().lower() for c in df.columns]
    
    if 'fecha' in df.columns:
        df['fecha'] = pd.to_datetime(df['fecha'], errors='coerce').dt.strftime('%Y-%m-%d')
        
    if 'cantidad' in df.columns:
        df['cantidad'] = df['cantidad'].apply(clean_val).astype(int)
    else:
        df['cantidad'] = 0
        
    if 'precio_unitario' in df.columns:
        df['precio_unitario'] = df['precio_unitario'].apply(clean_val)
    else:
        df['precio_unitario'] = 0.0
        
    if 'total_venta' in df.columns:
        df['total_venta'] = df['total_venta'].apply(clean_val)
    else:
        df['total_venta'] = 0.0
        
    # Recalcular si vino vacío por ser fórmula
    mask = (df['total_venta'] == 0)
    df.loc[mask, 'total_venta'] = df.loc[mask, 'cantidad'] * df.loc[mask, 'precio_unitario']
    
    return df

# --- 3. Subir y guardar datos preservando diseño, pestañas y fórmulas ---
def save_data_to_drive(file_id, df_to_save):
    """Actualiza el Excel conservando formato, estilos, anchos de columna y fórmulas."""
    service = get_drive_service()
    
    # Descargar el libro actual para no perder hojas adicionales (como Dashboard_Resumen)
    request = service.files().get_media(fileId=file_id, supportsAllDrives=True)
    fh = io.BytesIO()
    downloader = MediaIoBaseDownload(fh, request)
    done = False
    while not done:
        _, done = downloader.next_chunk()
    fh.seek(0)
    
    try:
        wb = openpyxl.load_workbook(fh)
    except Exception:
        wb = openpyxl.Workbook()

    # Seleccionar o crear la hoja Ventas
    if "Ventas" in wb.sheetnames:
        ws = wb["Ventas"]
        ws.delete_rows(1, ws.max_row + 10)  # Limpiar contenido anterior
    else:
        ws = wb.active
        ws.title = "Ventas"

    # Encabezados
    cols = ["id_transaccion", "fecha", "cliente", "ciudad", "categoria", "producto", "cantidad", "precio_unitario", "total_venta", "estado"]
    ws.append(cols)

    # Estilos del encabezado (Azul oscuro corporativo + texto blanco negrita)
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center")

    for col_idx in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment

    # Bordes sutiles para las filas
    thin_border = Border(
        left=Side(style='thin', color='E0E0E0'),
        right=Side(style='thin', color='E0E0E0'),
        top=Side(style='thin', color='E0E0E0'),
        bottom=Side(style='thin', color='E0E0E0')
    )

    # Insertar filas con fórmulas y formatos
    for row_idx, (_, row) in enumerate(df_to_save.iterrows(), start=2):
        cant = int(clean_val(row.get('cantidad', 0)))
        precio = float(clean_val(row.get('precio_unitario', 0.0)))
        formula_total = f"=G{row_idx}*H{row_idx}"  # Columna G (cantidad) * Columna H (precio)
        
        row_values = [
            str(row.get('id_transaccion', f"TRX-{row_idx-1:04d}")),
            str(row.get('fecha', '')),
            str(row.get('cliente', '')),
            str(row.get('ciudad', '')),
            str(row.get('categoria', '')),
            str(row.get('producto', '')),
            cant,
            precio,
            formula_total,
            str(row.get('estado', 'Completado'))
        ]
        ws.append(row_values)

        # Formato numérico y de moneda
        ws.cell(row=row_idx, column=7).number_format = '#,##0'
        ws.cell(row=row_idx, column=8).number_format = '$#,##0.00'
        ws.cell(row=row_idx, column=9).number_format = '$#,##0.00'

        for c_i in range(1, len(cols) + 1):
            c_cell = ws.cell(row=row_idx, column=c_i)
            c_cell.border = thin_border
            if c_i in (1, 2, 10):
                c_cell.alignment = Alignment(horizontal="center")

    # Ajuste automático del ancho de columnas para que no se encima nada
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(c.value or '')) for c in col)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Guardar en buffer y subir a Google Drive
    out_buf = io.BytesIO()
    wb.save(out_buf)
    out_buf.seek(0)

    media = MediaIoBaseUpload(
        out_buf,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        resumable=True
    )
    service.files().update(
        fileId=file_id,
        media_body=media,
        supportsAllDrives=True
    ).execute()

# --- 4. Encabezado de la App ---
col_title, col_btn = st.columns([5, 1])
with col_title:
    st.title("📈 Panel de Ventas en Vivo")
    st.caption("Conectado bidireccionalmente a Google Drive (base_datos_ventas.xlsx)")

with col_btn:
    st.write("")
    if st.button("🔄 Refrescar", help="Descarga los datos más recientes de Drive"):
        st.cache_data.clear()
        st.rerun()

FILE_ID = st.secrets["drive_settings"]["file_id"]

try:
    with st.spinner("Cargando datos desde Google Drive..."):
        df = load_data_from_drive(FILE_ID)
except Exception as e:
    st.error(f"Error al conectar con Google Drive: {e}")
    st.stop()

# --- 5. Pestañas: Dashboard Visual vs Editor Interactivo ---
tab_dash, tab_edit = st.tabs(["📊 Dashboard y Reportes", "✏️ Editor de Base de Datos"])

with tab_dash:
    st.sidebar.header("🔍 Filtros de Consulta")

    ciudades = ["Todas"] + sorted([c for c in df["ciudad"].dropna().unique().tolist() if str(c).strip()]) if "ciudad" in df.columns else ["Todas"]
    ciudad_sel = st.sidebar.selectbox("Ciudad:", ciudades)

    categorias = ["Todas"] + sorted([c for c in df["categoria"].dropna().unique().tolist() if str(c).strip()]) if "categoria" in df.columns else ["Todas"]
    cat_sel = st.sidebar.selectbox("Categoría:", categorias)

    estados = ["Todos"] + sorted([c for c in df["estado"].dropna().unique().tolist() if str(c).strip()]) if "estado" in df.columns else ["Todos"]
    estado_sel = st.sidebar.selectbox("Estado de orden:", estados)

    df_f = df.copy()
    if ciudad_sel != "Todas" and "ciudad" in df_f.columns:
        df_f = df_f[df_f["ciudad"] == ciudad_sel]
    if cat_sel != "Todas" and "categoria" in df_f.columns:
        df_f = df_f[df_f["categoria"] == cat_sel]
    if estado_sel != "Todos" and "estado" in df_f.columns:
        df_f = df_f[df_f["estado"] == estado_sel]

    # KPIs
    k1, k2, k3, k4 = st.columns(4)
    total_ventas = float(df_f["total_venta"].sum())
    total_unidades = int(df_f["cantidad"].sum())
    num_ordenes = len(df_f)
    ticket_medio = (total_ventas / num_ordenes) if num_ordenes > 0 else 0.0

    k1.metric("Ingresos Totales", f"${total_ventas:,.2f}")
    k2.metric("Unidades Vendidas", f"{total_unidades:,}")
    k3.metric("Ticket Promedio", f"${ticket_medio:,.2f}")
    k4.metric("Nº de Órdenes", num_ordenes)

    st.markdown("---")

    # Gráficos
    c_g1, c_g2 = st.columns(2)
    with c_g1:
        st.subheader("Ventas por Categoría")
        if "categoria" in df_f.columns and len(df_f) > 0:
            ventas_cat = df_f.groupby("categoria")["total_venta"].sum()
            st.bar_chart(ventas_cat)
        else:
            st.info("Sin datos para mostrar.")

    with c_g2:
        st.subheader("Ventas por Ciudad")
        if "ciudad" in df_f.columns and len(df_f) > 0:
            ventas_ciudad = df_f.groupby("ciudad")["total_venta"].sum()
            st.bar_chart(ventas_ciudad)
        else:
            st.info("Sin datos para mostrar.")

with tab_edit:
    st.subheader("📝 Edición directa en la Base de Datos")
    st.caption("Modifica celdas haciendo doble clic. La columna 'total_venta' se calcula sola en Excel mediante fórmula `=G*H`.")
    
    df_editado = st.data_editor(
        df,
        num_rows="dynamic",
        use_container_width=True,
        hide_index=True,
        disabled=["total_venta"]
    )
    
    if st.button("💾 Guardar cambios en Google Drive", type="primary"):
        try:
            with st.spinner("Guardando y formateando en Google Drive..."):
                save_data_to_drive(FILE_ID, df_editado)
            st.success("¡Base de datos actualizada con formato profesional y fórmulas en Google Drive!")
            st.cache_data.clear()
            st.rerun()
        except Exception as e:
            st.error(f"Error al guardar: {e}")
